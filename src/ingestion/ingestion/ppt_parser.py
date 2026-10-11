import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from pptx import Presentation

from src.ingestion.schemas.content_unit import ContentUnit
from src.ingestion.processing.chunker import chunk_text


def _extract_ppt_fallback_zip(file_path: str) -> list[tuple[int, str]]:
    """Extract slides text using zipfile XML parsing when python-pptx fails."""
    slides = []
    try:
        with zipfile.ZipFile(file_path, "r") as z:
            slide_names = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)]
            slide_names.sort(key=lambda s: int(re.search(r"\d+", s.split("/")[-1]).group()))
            for idx, name in enumerate(slide_names):
                xml_data = z.read(name)
                root = ET.fromstring(xml_data)
                texts = [elem.text for elem in root.iter() if elem.tag.endswith("}t") and elem.text and elem.text.strip()]
                full_text = " ".join(texts).strip()
                if full_text:
                    slides.append((idx + 1, full_text))
    except Exception as e:
        print(f"[PPT Parser] Zip XML extraction failed: {e}")
    return slides


def extract_ppt_content(file_path: str) -> list[ContentUnit]:
    source_name = Path(file_path).name
    source_stem = Path(source_name).stem
    content_units = []
    slide_data: list[tuple[int, str]] = []

    try:
        presentation = Presentation(file_path)
        for slide_index, slide in enumerate(presentation.slides, start=1):
            slide_text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        slide_text_parts.append(text)
            slide_text = "\n".join(slide_text_parts).strip()
            if slide_text:
                slide_data.append((slide_index, slide_text))
    except Exception as exc:
        print(f"[PPT Parser] python-pptx raised '{exc}', falling back to direct XML extraction...")
        slide_data = _extract_ppt_fallback_zip(file_path)

    if not slide_data:
        slide_data = _extract_ppt_fallback_zip(file_path)

    for slide_index, slide_text in slide_data:
        chunks = chunk_text(slide_text)
        for chunk_index, chunk in enumerate(chunks, start=1):
            unit = ContentUnit(
                unit_id=f"{source_stem}_s{slide_index}_c{chunk_index}",
                source_id=source_name,
                source_type="slides",
                location={"slide": slide_index},
                chunk_index=chunk_index,
                content_type="text",
                text=chunk,
                visual_description=None,
                image_path=None,
                topic=None,
                subtopic=None,
                concept_ids=[],
            )
            content_units.append(unit)

    return content_units