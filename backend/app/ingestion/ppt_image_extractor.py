from pathlib import Path
from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image


IMAGE_DIR = Path("extracted_images")
IMAGE_DIR.mkdir(exist_ok=True)


def extract_images_from_ppt(file_path: str) -> list[dict]:
    presentation = Presentation(file_path)

    source_name = Path(file_path).name
    source_stem = Path(source_name).stem

    extracted_images = []

    for slide_index, slide in enumerate(
        presentation.slides,
        start=1
    ):
        image_index = 0

        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue

            image_index += 1

            image = shape.image
            image_bytes = image.blob
            image_ext = image.ext.lower()

            # Convert GIF -> PNG using first frame
            if image_ext == "gif":
                try:
                    pil_image = Image.open(BytesIO(image_bytes))
                    pil_image.seek(0)

                    if pil_image.mode not in ("RGB", "RGBA"):
                        pil_image = pil_image.convert("RGBA")

                    image_name = (
                        f"{source_stem}_s{slide_index}_img{image_index}.png"
                    )

                    image_path = IMAGE_DIR / image_name

                    pil_image.save(
                        image_path,
                        format="PNG"
                    )

                    image_ext = "png"

                except Exception as exc:
                    print(
                        f"Failed to convert GIF on slide "
                        f"{slide_index}: {exc}"
                    )
                    continue

            else:
                image_name = (
                    f"{source_stem}_s{slide_index}_img{image_index}.{image_ext}"
                )

                image_path = IMAGE_DIR / image_name

                with open(image_path, "wb") as image_file:
                    image_file.write(image_bytes)

            # Optional dimension filter
            try:
                with Image.open(image_path) as img:
                    width, height = img.size
            except Exception:
                image_path.unlink(missing_ok=True)
                continue

            if width < 150 or height < 150:
                image_path.unlink(missing_ok=True)
                continue

            extracted_images.append(
                {
                    "source_id": source_name,
                    "slide": slide_index,
                    "image_index": image_index,
                    "image_path": str(image_path),
                    "extension": image_ext,
                    "width": width,
                    "height": height
                }
            )

    return extracted_images