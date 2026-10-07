from pathlib import Path

from pptx import Presentation

from app.schemas.content_unit import ContentUnit
from app.processing.chunker import chunk_text


def extract_ppt_content(file_path: str) -> list[ContentUnit]:
    presentation = Presentation(file_path)

    source_name = Path(file_path).name
    source_stem = Path(source_name).stem

    content_units = []

    for slide_index, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text_parts = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_text_parts.append(text)

        slide_text = "\n".join(slide_text_parts)

        if not slide_text.strip():
            continue

        chunks = chunk_text(slide_text)

        for chunk_index, chunk in enumerate(
            chunks,
            start=1
        ):
            unit = ContentUnit(
                unit_id=(
                    f"{source_stem}_s{slide_index}_c{chunk_index}"
                ),
                source_id=source_name,
                source_type="slides",
                location={
                    "slide": slide_index
                },
                chunk_index=chunk_index,
                content_type="text",
                text=chunk,
                visual_description=None,
                image_path=None,
                topic=None,
                subtopic=None,
                concept_ids=[]
            )

            content_units.append(unit)

    return content_units